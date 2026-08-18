# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from pathlib import Path

PHOTO = Path(r"C:\Users\masap\OneDrive\Documents\Lebenslauf-Foto-Cadima-Lusiola.png")

doc = Document()

for section in doc.sections:
    section.top_margin = Cm(1.4)
    section.bottom_margin = Cm(1.4)
    section.left_margin = Cm(1.7)
    section.right_margin = Cm(1.7)

NAVY = (30, 58, 95)
DARK = (25, 25, 25)
MUTED = (90, 90, 90)


def set_run_font(run, size=10, bold=False, color=None, name="Calibri"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def style_para(p, space_before=0, space_after=4, align=None):
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = 1.12
    return p


def add_para(text="", size=10, bold=False, color=None, space_before=0, space_after=4, align=None, container=None):
    parent = container if container is not None else doc
    p = parent.add_paragraph()
    style_para(p, space_before, space_after, align)
    if text:
        run = p.add_run(text)
        set_run_font(run, size=size, bold=bold, color=color)
    return p


def add_runs(parts, space_before=0, space_after=4, align=None, container=None):
    parent = container if container is not None else doc
    p = parent.add_paragraph()
    style_para(p, space_before, space_after, align)
    for text, size, bold, color in parts:
        run = p.add_run(text)
        set_run_font(run, size=size, bold=bold, color=color)
    return p


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "nil")
        tcBorders.append(element)
    tcPr.append(tcBorders)


def add_heading_bar(title):
    p = add_para(title.upper(), size=11, bold=True, color=NAVY, space_before=10, space_after=2)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), "1E3A5F")
    pBdr.append(bottom)
    pPr.append(pBdr)
    return p


def add_edu(years, title, place, detail=None):
    add_runs(
        [
            (years + "  ·  ", 10, False, MUTED),
            (title, 10, True, DARK),
        ],
        space_before=5,
        space_after=0,
    )
    add_para(place, size=10, color=(60, 60, 60), space_before=0, space_after=1)
    if detail:
        add_para(detail, size=9.5, color=(70, 70, 70), space_before=0, space_after=2)


# ===== HEADER with photo =====
header = doc.add_table(rows=1, cols=2)
header.autofit = True
header.allow_autofit = True
left, right = header.rows[0].cells
set_cell_border(left)
set_cell_border(right)
left.width = Cm(13.5)
right.width = Cm(3.8)
right.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

# Clear default paragraphs
left.paragraphs[0].clear()

add_para("CADIMA LUSIOLA", size=20, bold=True, color=NAVY, space_before=0, space_after=2, container=left)
add_para(
    "Informatiker Applikationsentwicklung  |  Junior Full-Stack Developer",
    size=10.5,
    color=(50, 70, 95),
    space_before=0,
    space_after=6,
    container=left,
)
add_para(
    "Höglerstrasse 55, 8600 Dübendorf",
    size=9.5,
    color=MUTED,
    space_before=0,
    space_after=1,
    container=left,
)
add_para(
    "+41 78 672 88 05  ·  clusiola4@gmail.com",
    size=9.5,
    color=MUTED,
    space_before=0,
    space_after=1,
    container=left,
)
add_para(
    "github.com/CapalotDaGreat  ·  www.cadimalusiola.com",
    size=9.5,
    color=MUTED,
    space_before=0,
    space_after=1,
    container=left,
)
add_para(
    "linkedin.com/in/cadima-lusiola-392833380",
    size=9.5,
    color=MUTED,
    space_before=0,
    space_after=0,
    container=left,
)

# Photo cell
rp = right.paragraphs[0]
rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run = rp.add_run()
run.add_picture(str(PHOTO), width=Cm(3.5))

# Profile
add_heading_bar("Profil")
add_para(
    "Motivierter Informatiker der Applikationsentwicklung (IMS) mit starkem Fokus auf moderne "
    "Webentwicklung. Ich entwickle klare, wartbare Anwendungen mit React, JavaScript/TypeScript, "
    "HTML und CSS und bringe eigene Full-Stack-Projekte von der Idee bis zur lauffähigen Demo. "
    "Zweisprachig aufgewachsen (Englisch / Deutsch), lernbereit und teamfähig – offen für eine "
    "Junior- oder Entry-Level-Stelle in der Softwareentwicklung.",
    size=10,
    color=DARK,
    space_before=4,
    space_after=2,
)

add_heading_bar("Personalien")
add_runs(
    [
        ("Geburtsdatum: ", 10, True, DARK),
        ("05.10.2005", 10, False, DARK),
        ("     ·     ", 10, False, MUTED),
        ("Nationalität: ", 10, True, DARK),
        ("Grossbritannien", 10, False, DARK),
        ("     ·     ", 10, False, MUTED),
        ("Vollständiger Name: ", 10, True, DARK),
        ("Cadima Mukasa Lusiola", 10, False, DARK),
    ],
    space_before=4,
    space_after=2,
)

add_heading_bar("Ausbildung")
add_edu(
    "2023 – 2026",
    "Informatikmittelschule (IMS)",
    "Kantonsschule Hottingen, Zürich",
    "Schwerpunkt Applikationsentwicklung: Softwareentwicklung, Webtechnologien, Datenbanken und praxisnahe Projektarbeit.",
)
add_edu(
    "2022 – 2023",
    "Gymnasium",
    "Kantonsschule Enge, Zürich",
    "Allgemeinbildende gymnasiale Ausbildung mit Fokus auf analytisches Denken und strukturiertes Lernen.",
)
add_edu("2021 – 2022", "10. Schuljahr", "Freie Katholische Schule Sumatra")
add_edu("2020 – 2021", "Sekundarschule", "Freie Katholische Schule Kreuzbühl")
add_edu("2018 – 2020", "Sekundarschule", "Schulhaus Grüze")
add_edu("2015 – 2018", "Primarschule", "Primarschule Högler")

add_heading_bar("Fachliche Kompetenzen")
add_runs(
    [
        ("Frontend: ", 10, True, DARK),
        ("React, JavaScript, HTML, CSS, responsive UI-Entwicklung", 10, False, DARK),
    ],
    space_before=4,
    space_after=2,
)
add_runs(
    [
        ("Backend & Daten: ", 10, True, DARK),
        ("Python, Java, MySQL, API-Integration, strukturierte Datenhaltung", 10, False, DARK),
    ],
    space_before=0,
    space_after=2,
)
add_runs(
    [
        ("Werkzeuge & Methoden: ", 10, True, DARK),
        ("Git / GitHub, Clean Code, iterative Entwicklung, Deployment von Webprojekten", 10, False, DARK),
    ],
    space_before=0,
    space_after=2,
)
add_runs(
    [
        ("Arbeitsweise: ", 10, True, DARK),
        (
            "lösungsorientiert, zuverlässig, schnelle Einarbeitung in neue Technologien, klare Kommunikation",
            10,
            False,
            DARK,
        ),
    ],
    space_before=0,
    space_after=2,
)

add_heading_bar("Ausgewählte Projekte")
add_runs(
    [
        ("Portfolio-Website  ·  ", 10, True, DARK),
        ("Next.js, React, TypeScript", 10, False, MUTED),
    ],
    space_before=4,
    space_after=0,
)
add_para(
    "Persönliche Portfolio-Plattform mit modernen UI-Komponenten, Projektübersicht und Kontaktformular – "
    "konzipiert und umgesetzt für professionelle Bewerbungspräsentation.",
    size=9.5,
    color=(70, 70, 70),
    space_before=0,
    space_after=3,
)
add_runs(
    [
        ("ScoutX & weitere Web-Apps  ·  ", 10, True, DARK),
        ("TypeScript, React, Node.js", 10, False, MUTED),
    ],
    space_before=0,
    space_after=0,
)
add_para(
    "Eigenständige Softwareprojekte mit Fokus auf saubere Architektur, Erweiterbarkeit und praxisnahe Features "
    "(u. a. Budget-Tracker, Tools und interaktive Anwendungen).",
    size=9.5,
    color=(70, 70, 70),
    space_before=0,
    space_after=3,
)
add_runs(
    [
        ("Code & Demos: ", 10, True, DARK),
        ("https://github.com/CapalotDaGreat  ·  http://www.cadimalusiola.com", 10, False, DARK),
    ],
    space_before=0,
    space_after=2,
)

add_heading_bar("Sprachkenntnisse")
add_runs(
    [
        ("Englisch – ", 10, True, DARK),
        ("Muttersprache", 10, False, DARK),
        ("     ·     ", 10, False, MUTED),
        ("Deutsch – ", 10, True, DARK),
        ("zweite Muttersprache", 10, False, DARK),
        ("     ·     ", 10, False, MUTED),
        ("Französisch – ", 10, True, DARK),
        ("B1", 10, False, DARK),
    ],
    space_before=4,
    space_after=2,
)

add_heading_bar("Interessen & Engagement")
add_para(
    "Aktives Mitglied beim FC Dübendorf; regelmässiges Fitness-Training. Teamgeist, Disziplin und "
    "Ausdauer aus dem Sport übertrage ich gezielt auf Projekte und Zusammenarbeit.",
    size=10,
    color=DARK,
    space_before=4,
    space_after=2,
)

add_heading_bar("Referenzen")
add_runs(
    [
        ("Tim Irmler", 10, True, DARK),
        ("  –  Informatiklehrer", 10, False, MUTED),
    ],
    space_before=4,
    space_after=0,
)
add_para(
    "tim.irmler@bzz.ch  ·  +41 79 840 33 80",
    size=9.5,
    color=(70, 70, 70),
    space_before=0,
    space_after=4,
)
add_runs(
    [
        ("Vivien Apafi", 10, True, DARK),
        ("  –  Wirtschaftslehrerin", 10, False, MUTED),
    ],
    space_before=0,
    space_after=0,
)
add_para(
    "vivien.apafi@ksh.ch  ·  +41 79 545 84 18",
    size=9.5,
    color=(70, 70, 70),
    space_before=0,
    space_after=2,
)

out1 = Path(r"c:\Users\masap\OneDrive\Documents\Lebenslauf-Cadima-Lusiola-Bewerbung.docx")
out1b = Path(r"c:\Users\masap\OneDrive\Documents\Lebenslauf-Cadima-Lusiola-Bewerbung-v2.docx")
out2 = Path(r"c:\Users\masap\Downloads\Cursed\my-portfolio\Lebenslauf-Cadima-Lusiola-Bewerbung.docx")
for path in (out1, out1b, out2):
    try:
        doc.save(path)
        print("Saved:", path)
    except PermissionError:
        print("Skipped (file open):", path)
